import openpyxl
from pathlib import Path
import pytest
from datetime import date

from kami_excel_extractor.heuristic_parser import HeuristicExcelParser
from kami_excel_extractor.core import KamiExcelExtractor
from kami_excel_extractor.schema import ExtractionOptions


def test_heuristic_parser_simple_table():
    parser = HeuristicExcelParser()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Simple"

    # Headers
    ws.append(["Name", "Age", "Department"])
    # Rows
    ws.append(["Alice", 30, "Engineering"])
    ws.append(["Bob", 25, "Design"])

    result = parser.parse_worksheet(ws)
    assert "data" in result
    assert len(result["data"]) == 2
    assert result["data"][0]["Name"] == "Alice"
    assert result["data"][0]["Age"] == 30
    assert result["data"][0]["Department"] == "Engineering"
    assert result["data"][1]["Name"] == "Bob"


def test_heuristic_parser_merged_headers():
    parser = HeuristicExcelParser()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Merged"

    # Setup complex header
    # Row 1: Merge A1:B1 for "Personal Details", C1 for "Status"
    ws.cell(row=1, column=1, value="Personal Details")
    ws.merge_cells("A1:B1")
    ws.cell(row=1, column=3, value="Status")

    # Row 2: Headers
    ws.cell(row=2, column=1, value="Name")
    ws.cell(row=2, column=2, value="Age")
    ws.cell(row=2, column=3, value="Active")

    # Row 3: Data
    ws.cell(row=3, column=1, value="Charlie")
    ws.cell(row=3, column=2, value=40)
    ws.cell(row=3, column=3, value="Yes")

    result = parser.parse_worksheet(ws)
    assert "data" in result
    assert len(result["data"]) == 1
    row_data = result["data"][0]
    # Check propagated header keys
    assert "Personal Details_Name" in row_data
    assert "Personal Details_Age" in row_data
    assert "Status_Active" in row_data
    assert row_data["Personal Details_Name"] == "Charlie"
    assert row_data["Personal Details_Age"] == 40
    assert row_data["Status_Active"] == "Yes"


def test_heuristic_parser_key_value_form():
    parser = HeuristicExcelParser()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Form"

    # Row 1: Label and Value
    ws.cell(row=1, column=1, value="Project Name:")
    ws.cell(row=1, column=2, value="Kami Extractor")

    # Row 2: Label and Merged Value
    ws.cell(row=2, column=1, value="Description：")
    ws.cell(row=2, column=2, value="A great rule-based parser")
    ws.merge_cells("B2:D2")

    # Row 4: Merged Label and value below it
    ws.cell(row=4, column=1, value="Manager Name")
    ws.merge_cells("A4:B4")
    ws.cell(row=5, column=1, value="Jules")

    result = parser.parse_worksheet(ws)
    assert "data" in result
    form_data = result["data"]
    assert form_data["Project Name"] == "Kami Extractor"
    assert form_data["Description"] == "A great rule-based parser"
    assert form_data["Manager Name"] == "Jules"


@pytest.mark.asyncio
async def test_kami_excel_extractor_heuristic_integration(tmp_path):
    # Save a temporary workbook
    excel_path = tmp_path / "test_heuristic.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project"
    ws.append(["Task", "Status"])
    ws.append(["Task 1", "Done"])
    ws.append(["Task 2", "In Progress"])
    wb.save(excel_path)

    extractor = KamiExcelExtractor(output_dir=tmp_path)
    options = ExtractionOptions(model="heuristic")
    result = await extractor.aextract_structured_data(excel_path, options=options)

    assert "sheets" in result
    assert "Project" in result["sheets"]
    sheet_data = result["sheets"]["Project"]
    assert "data" in sheet_data
    assert len(sheet_data["data"]) == 2
    assert sheet_data["data"][0]["Task"] == "Task 1"
    assert sheet_data["data"][0]["Status"] == "Done"
